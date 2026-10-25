"""Build the editable Bramha Atlas pitch. Run with python-pptx + Pillow installed.
All measurements originate in the copied public reference report; diagrams are editable.
"""
from pathlib import Path
import json, shutil, hashlib, math, textwrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE, XL_LABEL_POSITION, XL_TICK_MARK, XL_LEGEND_POSITION
from PIL import Image
import diagrams_3d as d3

OUT=Path(__file__).resolve().parents[1]
REPO=OUT.parent
ASSETS=OUT/'assets'
SOURCES=OUT/'source'
for dest,src in {
 'reference-check.json':'artifacts/experiment-analysis/real-reference-check.json',
 'mapping-audit.json':'data/manifests/pten-mapping-audit.json',
 'agent-run.json':'artifacts/experiment-analysis/live-team-verified-report.json',
 'storage-check.json':'artifacts/data-console/actual-storage-summary.json',
 'argon-check.json':'artifacts/ARGON_CONNECTIVITY.json',
 'assay-dictionary.json':'data/manifests/pten-assay-dictionary.json',
 'source-manifest.json':'data/manifests/pten-source-manifest.json',
}.items():
 if (REPO/src).exists():shutil.copy2(REPO/src,SOURCES/dest)
for dest,src in {
 'emblem.png':'frontend/public/branding/bramha-emblem.png',
 'linked-structure-results.png':'artifacts/experiment-analysis/actual-linked-structure-results.png',
 'agent-network.png':'artifacts/experiment-analysis/actual-autonomous-team-network.png',
 'world-evidence.png':'artifacts/experiment-analysis/world-map-adjacent-results-light.png',
}.items():
 if (REPO/src).exists():shutil.copy2(REPO/src,ASSETS/dest)
REF=json.loads((SOURCES/'reference-check.json').read_text())
TEAM=json.loads((SOURCES/'agent-run.json').read_text())
STORE=json.loads((SOURCES/'storage-check.json').read_text())
rows={v['protein_hgvs'].removeprefix('p.'):v for v in REF['variants']}
BG='F6F3EB';PAPER='FFFDF8';INK='193B3B';ACC='226B60';SEA='A5DFD0';COP='9E5837';MUT='58706F';LINE='C8D4CD';DARK='0B1923';WHITE='F2E9D7'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='Bramha Atlas — From scattered evidence to a reviewable research decision'
prs.core_properties.subject='PTEN research prototype, verified artifacts and design-partner proposal'
prs.core_properties.author='Bramha Atlas'
prs.core_properties.keywords='PTEN, research, evidence, computational analysis, pitch'
FONT='Arial'
scene=[];notes=[];ledger=[];asset_uses=[]
def rgb(c):return RGBColor.from_string(c)
def rect(s,x,y,w,h,fill=PAPER,line=None,radius=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 sh.line.fill.background() if not line else None
 if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(.8)
 if radius:
  try:sh.adjustments[0]=.06
  except:pass
 scene[-1]['objects'].append(dict(kind='rect',x=x,y=y,w=w,h=h,fill=fill,line=line))
 return sh

def txt(s,text,x,y,w,h,size=20,color=INK,bold=False,font=FONT,align=None,link=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(4);p.space_before=Pt(0);p.line_spacing=1.12
  if align is not None:p.alignment=align
  if link:
   for r in p.runs:r.hyperlink.address=link
 scene[-1]['objects'].append(dict(kind='text',text=text,x=x,y=y,w=w,h=h,size=size,color=color,bold=bold))
 return sh

def line(s,x1,y1,x2,y2,color=LINE,width=1.2):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
 scene[-1]['objects'].append(dict(kind='line',x=x1,y=y1,x2=x2,y2=y2,color=color,width=width));return sh

def image(s,name,x,y,w,h,crop=None):
 p=ASSETS/name
 with Image.open(p) as im: iw,ih=im.size
 if crop:
  l,t,r,b=crop;sw=(r-l)*iw;shh=(b-t)*ih
  pic=s.shapes.add_picture(str(p),Inches(x),Inches(y),width=Inches(w),height=Inches(h));pic.crop_left=l;pic.crop_top=t;pic.crop_right=1-r;pic.crop_bottom=1-b
 else:
  scale=min(w/iw,h/ih);ww=iw*scale;hh=ih*scale;x+=(w-ww)/2;y+=(h-hh)/2;w=ww;h=hh
  pic=s.shapes.add_picture(str(p),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
 asset_uses.append(dict(slide=len(scene),asset=name,crop=crop))
 scene[-1]['objects'].append(dict(kind='image',asset=name,x=x,y=y,w=w,h=h,crop=crop));return pic

def slide(title,section,source,note,dark=False):
 s=prs.slides.add_slide(prs.slide_layouts[6]);scene.append(dict(title=title,dark=dark,objects=[]));bg=s.background.fill;bg.solid();bg.fore_color.rgb=rgb(DARK if dark else BG)
 txt(s,section.upper(),.55,.32,11,.3,10,SEA if dark else ACC,True)
 txt(s,title,.55,.87,12.2,1.05,32,WHITE if dark else INK,True)
 line(s,.55,7.01,12.78,7.01,'31515C' if dark else LINE,.7)
 txt(s,'BRAMHA ATLAS  /  25 OCTOBER 2026',.55,7.12,8,.17,9,'B6C5C7' if dark else MUT)
 txt(s,f'{len(scene):02d}',12.15,7.10,.6,.22,10,SEA if dark else ACC,True,align=PP_ALIGN.RIGHT)
 s.notes_slide.notes_text_frame.text=note+'\n\nSOURCES / EVIDENCE\n'+source+'\n\nEvidence reviewed: 25 October 2026. Saved artifacts are historical observations, not a fresh live-provider certification.'
 notes.append(dict(slide=len(scene),title=title,notes=note,sources=source))
 ledger.append(dict(slide=len(scene),title=title,sources=source))
 return s

def pill(s,label,x,y,w,fill=ACC,color=PAPER):
 rect(s,x,y,w,.36,fill,radius=True);txt(s,label,x+.12,y+.065,w-.24,.24,11,color,True)
def card(s,x,y,w,h,label,body,number=None):
 rect(s,x,y,w,h,PAPER,LINE,True)
 if number:txt(s,number,x+.2,y+.2,w-.4,.5,30,COP,True);yy=y+.92
 else:yy=y+.23
 txt(s,label,x+.2,yy,w-.4,.5,22,INK,True);txt(s,body,x+.2,yy+.65,w-.4,h-(yy-y)-.75,18,MUT)
def table(s,headers,data,x,y,width,colwidths=None,row_h=.55,font=17):
 colwidths=colwidths or [width/len(headers)]*len(headers)
 for ri,row in enumerate([headers]+data):
  xx=x
  rect(s,x,y+ri*row_h,width,row_h,INK if ri==0 else (PAPER if ri%2 else 'EBEFE7'))
  for ci,cell in enumerate(row):
   txt(s,str(cell),xx+.14,y+ri*row_h+.13,colwidths[ci]-.25,row_h-.12,font,WHITE if ri==0 else INK,ri==0);xx+=colwidths[ci]
 return y+(len(data)+1)*row_h

BASE='http://localhost:8500/research/'+REF['project_id']+'/workspace?view=analysis&analysis_run='+REF['baseline_run_id']
FOLLOW='http://localhost:8500/research/'+REF['project_id']+'/workspace?view=analysis&analysis_run='+REF['followup_run_id']
RUN='http://localhost:8500/research/'+REF['project_id']+'/investigations/'+TEAM['run_id']
EVID='source/reference-check.json; source/assay-dictionary.json; source/source-manifest.json. Original records: https://www.mavedb.org/score-sets/urn:mavedb:00000054-a-1 and https://www.mavedb.org/score-sets/urn:mavedb:00000102-0-1. Dataset '+REF['dataset_version']+'.'
# 1 — editorial cover
s=slide('Bramha Atlas','Evidence that leads somewhere','Production emblem: frontend/public/branding/bramha-emblem.png; source/reference-check.json; assets/linked-structure-results.png.',
'''OPEN (20 seconds). Which mutation should we investigate next, and why? Bramha Atlas connects existing measurements, source evidence and reference structure into a research decision a scientist can inspect and reproduce. This is a computational research prototype. The frame is an actual saved PTEN workspace, not an animated simulation.\nOpening demo: use the button to My Research and select “Inspect saved baseline”. Walk through the 90-second script in DEMO.md. Recorded evidence is the reliable default; do not call it a live Argon run. If the local app is unavailable, use the supplied frames and continue on slide 4.''')
# cover title uses consistent header, visual dominates
image(s,'emblem.png',.64,2.15,1.15,1.65)
txt(s,'Which mutation should\nwe investigate next,\nand why?',.65,3.95,6,1.75,30,INK,True)
txt(s,'Reviewable research, connected to evidence.',.65,5.91,6,.4,18,MUT)
pill(s,'OPEN THE PTEN DEMO  →',.65,6.43,2.9)
# actual clickable button hyperlink
s.shapes[-2].click_action.hyperlink.address='http://localhost:8500/research'
rect(s,7.0,1.95,5.78,4.82,DARK,radius=True)
image(s,'linked-structure-results.png',7.15,2.10,5.48,4.39,crop=(.016,.229,.476,.699))
txt(s,'ACTUAL REFERENCE WORKSPACE  /  PTEN · 1D5R',7.15,6.52,5.5,.18,9,SEA,True)
# 2
s=slide('The next experiment starts with an evidence decision','The problem','Customer scenario is an illustrative workflow hypothesis; no measured time savings or customer traction claimed.',
'''25 seconds. Imagine a small functional-genomics team with several variant explanations and limited experimental capacity. A public activity score, an abundance observation, a structure and a paper each tell only part of the story. The working hypothesis is that the expensive handoff is connecting those sources to a defensible next question. We have not measured the frequency or cost of this pain with customers yet.''')
txt(s,'One researcher.\nCompeting explanations.\nLimited lab capacity.',.65,2.1,5.2,2.5,31,INK,True)
for y,label,body in [(2.1,'MEASUREMENTS','What did each assay measure?'),(3.3,'STRUCTURE','Can this position be mapped?'),(4.5,'LITERATURE','Does the cited passage apply?')]:
 rect(s,7,y,5.67,.92,PAPER,LINE,True);txt(s,label,7.22,y+.14,5,.2,10,ACC,True);txt(s,body,7.22,y+.43,5.15,.35,20)
rect(s,.65,6.05,12.02,.57,INK);txt(s,'The proposed job: assemble a research decision with evidence and explicit unknowns.',.87,6.18,11.6,.33,20,WHITE)
# 3
s=slide('From a question to a brief you can check','The product','docs/EXPERIMENT_EXECUTION_AUDIT.md; docs/EXPERIMENT_ANALYSIS_DELIVERY.md; backend/app/analysis_execution.py; backend/app/analysis_results.py.',
'''25 seconds. The unit of value is a reviewable investigation, rather than an isolated chat response. Data and numerical operations retain inputs and versions. The researcher checks the result and chooses a follow-up. The diagram describes the supported investigation workflow; it does not mean every step invokes a model. The quantitative mission runs deterministic Python without an LLM.''')
d3.workflow(s,txt,line,rect)
# 4
s=slide('Three real substitutions. One bounded question.','The PTEN demonstration',EVID,
'''40 seconds. Open the saved baseline, not the model execution screen. The three selected substitutions are C124S, R172H and Y16T. These are human PTEN protein substitutions against canonical P60484, not uniquely inferred DNA alleles. The snapshot combines two distinct public MaveDB score sets and accepted reference mappings. All three selected positions map; this case does not demonstrate missing coordinates. The missingness is R172H abundance. Ask: which selected variants meet the assay-specific criteria, what structural context is available, and what remains unresolved? Show the saved inputs and the missing row.\nLocal baseline: '''+BASE)
for i,(v,label,body) in enumerate([('C124S','Paired measurements','Activity and abundance\nare both available.'),('R172H','A missing observation','Activity is available;\nabundance is missing.'),('Y16T','A different profile','Paired measurements;\nexcluded by baseline rule.')]):
 x=.65+i*4.08;rect(s,x,2.05,3.83,2.55,PAPER,LINE,True);txt(s,v,x+.24,2.3,3.35,.6,36,COP,True);txt(s,label,x+.24,3.08,3.38,.4,21,INK,True);txt(s,body,x+.24,3.65,3.3,.73,18,MUT)
if (ASSETS/'demo-baseline-controls.png').exists():
 image(s,'demo-baseline-controls.png',.65,4.94,7.1,1.8,crop=(.169,.655,.981,.996))
else:txt(s,'PINNED DATA\nHuman PTEN · P60484\nExperimental reference · 1D5R',.65,5.08,7,1.3,22)
txt(s,'Which profiles meet our rule?\nWhat remains unresolved?',8.15,5.1,4.5,1.35,24,INK,True)
# 5 native chart
s=slide('Two assays answer different questions','Measurements, with their meaning',EVID,
'''40 seconds. The x axis is published enrichment/fitness in a humanized yeast model; it is not a direct clinical activity score. The y axis is the native VAMP-seq abundance score, not a probability and not clamped to 0–1. The chart contains only the two paired profiles. R172H remains visible in the accompanying table, but is not plotted with an invented y value. Values on the slide are rounded for readability; the workbook and reference JSON preserve full precision. The displayed boundaries are exploratory analytical rules. Native activity/abundance values cannot be subtracted or treated as a common scale. Reported uncertainty metadata is retained in the source; no inferential conclusion is drawn here.''')
cd=XyChartData()
for name in ['C124S','Y16T']:
 series=cd.add_series(name);v=rows[name];series.add_data_point(v['activity'],v['abundance'])
gf=s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER,Inches(.65),Inches(2.05),Inches(7),Inches(4.3),cd);chart=gf.chart
chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.has_title=False
chart.chart_style=13
chart.category_axis.crosses_at=-5.5;chart.value_axis.crosses_at=.4
chart.category_axis.major_unit=1;chart.value_axis.major_unit=.2
chart.value_axis.has_major_gridlines=False
chart.font.name=FONT;chart.font.size=Pt(13)
for axis,lo,hi,title in [(chart.category_axis,-5.5,.5,'Activity · native yeast enrichment'),(chart.value_axis,.4,1.2,'Abundance · native VAMP-seq')]:
 axis.minimum_scale=lo;axis.maximum_scale=hi;axis.has_title=True;axis.axis_title.text_frame.text=title
 axis.axis_title.text_frame.paragraphs[0].font.size=Pt(14);axis.axis_title.text_frame.paragraphs[0].font.name=FONT
 axis.tick_labels.font.size=Pt(12);axis.tick_labels.font.name=FONT;axis.major_tick_mark=XL_TICK_MARK.NONE
for i,ser in enumerate(chart.series):
 ser.marker.style=XL_MARKER_STYLE.CIRCLE if i==0 else XL_MARKER_STYLE.DIAMOND;ser.marker.size=12
 ser.marker.format.fill.solid();ser.marker.format.fill.fore_color.rgb=rgb(ACC if i==0 else COP);ser.marker.format.line.color.rgb=rgb(ACC if i==0 else COP)

scene[-1]['objects'].append(dict(kind='chart',x=.65,y=2.05,w=7,h=4.3))
table(s,['Variant','Activity','Abundance'],[[n,f"{rows[n]['activity']:.3f}",'Missing' if rows[n]['abundance'] is None else f"{rows[n]['abundance']:.3f}"] for n in ['C124S','R172H','Y16T']],8.05,2.3,4.62,[1.55,1.5,1.57],row_h=.66,font=16)
txt(s,'2 of 3 have both values.\nMissing stays unknown.',8.05,5.33,4.6,1.1,22,INK,True)
txt(s,'Public MaveDB measurements · pinned PTEN snapshot · no cross-assay normalization',.85,6.6,11.8,.2,11,MUT)
#6
s=slide('A mapped position gives context, not causation','Reference structure',EVID+' Reference coordinates: https://www.rcsb.org/structure/1D5R ; UniProt: https://www.uniprot.org/uniprotkb/P60484/entry .',
'''30 seconds. These are experimental reference coordinates from 1D5R. The mapping checks sequence identity and observed C-alpha correspondence. It is a precomputed audited reference-to-construct alignment, not a new alignment every time the viewer opens. All three chosen substitutions are mapped in this saved report. A selected side chain or colored residue is a substitution annotation; it does not compute a mutant fold. Distances are reference-coordinate Euclidean distances when coordinates are available; proximity alone does not establish a shared mechanism.''')
molecular_diagram=d3.molecule(s,txt,line,rect,SOURCES)
#7
s=slide('Change one rule. See exactly what changes.','A reproducible follow-up',EVID+' Baseline '+REF['baseline_run_id']+'; follow-up '+REF['followup_run_id']+'.',
'''40 seconds. This is a saved analytical sensitivity check, separate from the model-backed investigation. The activity ceiling changes from -1.11 to 10. The abundance boundary remains 0.6121866579147568. Selection, source versions, mapping policy and the paired-measurement requirement remain fixed. A ceiling of 10 is deliberately permissive, not an established biological threshold. C124S remains selected and Y16T enters. R172H stays unresolved because its abundance value is absent. This changed analytical set does not demonstrate a new biological effect, a causal mechanism or an agent-generated hypothesis revision. Baseline and follow-up made zero model requests.\nSaved follow-up: '''+FOLLOW)
for x,label,value,body in [(.65,'BASELINE','C124S','Activity ceiling −1.11'),(7.0,'FOLLOW-UP','C124S + Y16T','Activity ceiling 10')]:
 rect(s,x,2.05,5.68,2.42,PAPER,LINE,True);txt(s,label,x+.24,2.3,5.1,.25,11,ACC,True);txt(s,value,x+.24,2.9,5.1,.65,32,INK,True);txt(s,body,x+.24,3.83,5.1,.35,20,MUT)
txt(s,'→',6.48,2.95,.45,.55,30,COP,True)
for x,label,body in [(.65,'ENTERED','Y16T'),(4.73,'LEFT','None'),(8.81,'UNCHANGED','C124S')]:
 txt(s,label,x,4.95,3.7,.25,11,ACC,True);txt(s,body,x,5.43,3.7,.6,30,INK,True)
pill(s,'EXPLORATORY CRITERION CHANGE · 0 MODEL REQUESTS',.65,6.32,6.6,INK)
#8
s=slide('Scoped reviews, inspectable handoffs','Recorded agent investigation','source/agent-run.json; backend/app/orchestration/agent_network.py; assets/agent-network.png. Actual provider: OpenRouter / x-ai/grok-4.3. Argon remains unverified; source/argon-check.json.',
'''45 seconds. This diagram reflects one actual saved completed-with-limitations run. There is one model research lead and four delegated workers: a deterministic identity validator, an independent model assay reviewer, an independent model structure reviewer, and a deterministic artifact verifier. Both independent model reviews began after identity validation and their saved intervals overlap. There were five model requests, ten deterministic tool calls and eleven saved messages. This proves this saved task network and artifact handoffs, not external A2A protocol interoperability or unrestricted replication. It used OpenRouter / x-ai/grok-4.3; do not attribute this run to Argon. Argon is the intended integration and remains unverified. No private reasoning traces are displayed.\nOpen saved recorded result: '''+RUN)
d3.agents(s,txt,line,rect)
#9
s=slide('The useful result includes the unknowns','The research outcome',EVID+' source/agent-run.json. Adaptive limitation-driven model revision is not verified by this threshold comparison.',
'''35 seconds. The baseline supports a descriptive statement: C124S meets this declared paired rule. It does not establish the biological reason. R172H cannot be evaluated with that paired rule because abundance is missing. The follow-up adds Y16T only after the criterion is changed. The next research question is whether additional applicable evidence can distinguish functional loss from abundance-related explanations. That is a research proposal, not a reported result or causal conclusion. A genuine autonomous limitation-driven revision is not established by this artifact set; the deterministic criterion change must not be sold as that behavior.''')
rect(s,.65,2.1,7.4,3.3,INK,radius=True)
txt(s,'SUPPORTED BY THE BASELINE',.9,2.38,6.8,.3,11,SEA,True)
txt(s,'C124S meets the declared\npaired-measurement rule.',.9,3.0,6.75,1.25,31,WHITE,True)
txt(s,'A descriptive shortlist for scientific review.',.9,4.62,6.8,.5,20,SEA)
txt(s,'Still unresolved',8.6,2.2,4,.5,26,INK,True)
txt(s,'R172H: abundance is missing.\n\nMechanism and disease relevance\nare not established here.',8.6,3.0,4.05,2.1,21,MUT)
line(s,.65,5.85,12.67,5.85)
txt(s,'NEXT QUESTION',.65,6.1,2.1,.3,11,ACC,True)
txt(s,'What additional applicable evidence could distinguish the competing explanations?',3.0,6.02,9.65,.7,23,INK,True)
#10
s=slide('Connected evidence, explicit storage boundaries','System architecture','docs/MONGODB_DATA_CONSOLE.md; source/storage-check.json; backend/app/analysis_execution.py; backend/app/corpus_store.py; backend/app/research_hub.py; frontend/src/workspace/ResearchHub.tsx.',
'''35 seconds. The public corpus is persisted in MongoDB Atlas; project questions, investigations and operation results remain in the configured local store. Deterministic Python tools execute numerical analysis. The lead and scoped model reviewers interpret bounded evidence, with artifact checks. The contextual Bramha Assistant is distinct from investigation agents. The stored passage vectors use a local pretrained multilingual embedding model. A recorded ready Atlas vector index does not by itself certify a full hybrid retrieval query; that integration must be separately verified. This is a conceptual architecture diagram, not a claim that every source is always queried for every run.''')
d3.architecture(s,txt,line,rect)
#11
s=slide('A prototype with evidence you can inspect','Demonstrated scope','source/reference-check.json; source/agent-run.json; source/storage-check.json; source/argon-check.json; docs/EXPERIMENT_ANALYSIS_DELIVERY.md. Counts are the recorded 25 October 2026 snapshot.',
'''35 seconds. These rows distinguish demonstrated and pending work. The numerical mission has an audited saved baseline and follow-up. The independent agent team has a successful saved fallback-provider run, completed with limitations. Atlas has a connected public-corpus snapshot with 480 studies and 4,784 passages, distinct from project-state persistence. Supplied-model alignment also exists but is not the core pitch demonstration. Argon scientific execution, a verified end-to-end hybrid retrieval claim, autonomous limitation-driven revision in this case, and clinical validation are not established by these artifacts. Repository software checks do not validate biology.''')
table(s,['Capability','Evidence in hand','Status'],[
 ['PTEN analysis','Pinned baseline + one-factor follow-up','Demonstrated'],
 ['Independent reviews','Saved model-backed team run','Recorded; with limits'],
 ['Public corpus','Atlas connection + stored records','Snapshot verified'],
 ['Argon scientific run','Connectivity attempt timed out','Unverified'],
 ['Clinical or causal validity','No qualifying validation in this case','Not established'],
],.65,2.08,12.02,[3.15,5.35,3.52],row_h=.6,font=17)
txt(s,'480 studies  /  4,784 passages',.65,6.1,7.5,.55,28,INK,True)
txt(s,'Recorded Atlas snapshot · 25 Oct 2026, 05:47 UTC',8.0,6.18,4.7,.65,14,MUT)
#12
s=slide('Start with teams making variant research decisions','Initial customer','Proposed customer segmentation and pilot design are hypotheses; no signed pilots or validated willingness to pay claimed.',
'''30 seconds. Start with academic functional-genomics labs and small biotech research teams already interpreting public or internally reviewed variant evidence. The first job is a reviewable computational shortlist and an explicit list of missing evidence. The proposal is a small design-partner pilot using scoped, nonclinical research questions. We have not established customer demand, procurement fit or willingness to pay. Use expert feedback to narrow the workflow before expanding genes or organization-wide use.''')
card(s,.65,2.15,5.78,2.55,'Academic genomics labs','Review variant evidence before\nchoosing a follow-up experiment.','01')
card(s,6.88,2.15,5.79,2.55,'Small biotech teams','Connect measurements, structure\nand source review in one record.','02')
rect(s,.65,5.3,12.02,1.17,INK,radius=True);txt(s,'FIRST PILOT JOB',.9,5.53,2.3,.25,11,SEA,True);txt(s,'A reviewed shortlist and a traceable next research question.',3.1,5.6,9.25,.65,24,WHITE,True)
#13
s=slide('A focused workflow in an established ecosystem','Competitive context',
'''Official pages checked 25 October 2026: Benchling https://www.benchling.com/ ; Galaxy https://galaxyproject.org/ ; MaveDB https://mavedb.org/docs/mavedb/index.html and https://www.mavedb.org/docs/mavedb/mavemd/variant-page.html ; Europe PMC https://europepmc.org/about . Market context: Grand View Research https://www.grandviewresearch.com/horizon/outlook/drug-discovery-informatics-market-size/global .''',
'''40 seconds. These products overlap with parts of the workflow and should be treated fairly. Benchling describes agentic R&D workflows and scientific context, so do not claim that competitors lack agents. Galaxy provides reproducible computational workflows. MaveDB provides variant-effect data; its related interfaces already expose multi-assay and structural interpretation. Europe PMC enables life-science literature discovery and reading. Our proposed distinction is a narrowly scoped investigation record that connects declared criteria, actual tool artifacts, source coverage and explicit unknowns; uniqueness and superiority are not established. Grand View Research estimates the broad global drug-discovery-informatics category at $3.65 billion in 2024. That is market context, not Bramha Atlas TAM, valuation, attainable revenue or a primary-source measurement. We do not sum overlapping markets.''')
table(s,['Existing tool','Established focus'],[
 ['Benchling','R&D workflows, context and agents'],['Galaxy','Reproducible computational workflows'],['MaveDB','Variant-effect datasets and interpretation'],['Europe PMC','Life-science literature discovery'],
],.65,2.05,7.8,[2.05,5.75],row_h=.7,font=18)
txt(s,'$3.65B',9.02,2.3,3.5,.8,44,INK,True);txt(s,'2024 global drug-discovery\ninformatics estimate',9.02,3.28,3.6,1.0,19,MUT);txt(s,'Broad category context\nGrand View Research\nNot our addressable market',9.02,4.7,3.6,1.15,15,MUT)
rect(s,.65,6.0,12.02,.65,INK);txt(s,'Our proposed focus: one inspectable decision, from declared criteria to unresolved evidence.',.85,6.16,11.6,.35,19,WHITE)
# Market gap, customer economics and a bounded entry segment.
MARKET_SOURCES='Official sources checked 25 October 2026: Benchling academic plan and AI https://www.benchling.com/academic ; Galaxy https://galaxyproject.org/ ; MaveDB variant interface https://mavedb.org/docs/mavedb/mavemd/variant-page.html . Competitor coverage does not prove unmet demand. Market opportunity and value assumptions: source/market-model.json; Market-Value-Model.xlsx.'
s=slide('The gap is the decision handoff','Market gap to validate',MARKET_SOURCES,
'''35 seconds. The proposed gap is the work between finding evidence and approving a research decision: checking that records refer to the same substitution, respecting distinct assay meanings, linking structural coverage, explaining a changed shortlist, and preserving unresolved evidence. Existing tools overlap substantially. Benchling already has academic tools, AI and source-linked insights; Galaxy has reproducible workflows; MaveDB has assay and variant interpretation. We cannot claim those tools cannot solve the problem or that Bramha is unique. Our wedge is a preconfigured, bounded variant investigation with inspectable calculations and a reviewable outcome. Customer interviews and a task-level comparison must establish whether this handoff still costs enough time and whether the prototype reduces it. Start with small research teams that lack a dedicated bioinformatics workflow; that segment is a hypothesis. No signed pilots or validated demand are implied.''')
for i,(label,body) in enumerate([('Evidence exists','Assay records, structures\nand relevant passages.'),('The proposed gap','A scientist still needs a\ndefensible next decision.'),('Bramha’s wedge','A checked shortlist,\nchanged inputs and unknowns.')]):
 x=.65+i*4.08;d3.solid(s,x,3.7,3.5,.57,.25,-.22,INK if i==1 else ACC)
 txt(s,label,x+.18,3.83,3.14,.42,20,WHITE,True);txt(s,body,x,4.65,3.8,1.0,20,MUT)
 if i<2:txt(s,'→',x+3.7,3.75,.35,.4,23,COP,True)
txt(s,'Why pay when capable free tools exist?',.65,2.1,12,.55,28,INK,True)
txt(s,'Only if this specific handoff becomes measurably easier and still withstands expert review.',.65,6.12,12,.65,22,ACC,True)
s=slide('What is closing that gap worth to one lab?','Customer value — illustrative',MARKET_SOURCES,
'''40 seconds. All inputs here are assumptions, not measured savings. Suppose a lab handles eight applicable investigations per month and Bramha recovers two net researcher hours per investigation after additional review, onboarding and correction time. At an assumed loaded research-capacity value of sixty US dollars per hour, annual capacity value is 8 times 2 times 60 times 12, or $11,520. At the proposed $300 monthly lab fee, the annual fee is $3,600, leaving $7,920 of illustrative capacity value after the subscription alone. This is released capacity, not guaranteed payroll savings, profit, or a measured ROI. Benefits are 3.2 times the fee before any other unmodelled costs. Break-even is five net hours recovered per month at this hourly value. A low-use lab with only two investigations monthly at two hours per case produces $2,880 annual capacity value, below the fee; it should not buy on this assumption. Do not attach a monetary value to avoided wet-lab failures or better scientific decisions without evidence. The pilot must measure end-to-end time, reviewer effort, errors and repeated usage against the current workflow.''')
rect(s,.65,2.1,7.5,2.25,INK,radius=True)
txt(s,'8 cases/month × 2 net hours × $60/hour',.94,2.43,6.9,.6,25,WHITE,True)
txt(s,'= $11,520/year',.94,3.28,6.9,.65,36,SEA,True)
txt(s,'ILLUSTRATIVE RESEARCH CAPACITY VALUE',.94,4.0,6.9,.2,10,SEA,True)
txt(s,'5 net hours/month',8.6,2.35,4.1,.95,29,INK,True)
txt(s,'to cover a $300/month fee\nat $60 per researcher hour.',8.6,3.57,4.05,1.0,20,MUT)
for x,head,body in [(.65,'Annual lab fee','$3,600 · proposed'),(4.73,'Value after fee','$7,920 · scenario'),(8.81,'Low use can fail','2 cases/month: below fee')]:
 txt(s,head,x,5.12,3.8,.4,21,ACC,True);txt(s,body,x,5.76,3.8,.72,20,INK)
txt(s,'Assumed net time recovered; capacity value is not measured cash savings.',.65,6.65,12,.23,11,MUT)
s=slide('A testable US entry segment: $0.72–2.16M/year','Bottom-up opportunity — scenario',
'BIO, America’s Innovation Engine: The Power of Small and Mid-Sized Biotechs, January 2026: https://www.bio.org/sites/default/files/2026-01/the_power_of_small_and_mid-size_biotechs.pdf . BIO states more than 2,000 US biotech companies. This model uses a rounded 2,000-company reference pool, not a count of qualified buyers. All fit, account and price inputs are assumptions: source/market-model.json; Market-Value-Model.xlsx.',
'''45 seconds. BIO’s January 2026 industry brief describes more than two thousand biotech companies in the United States. We use a rounded two-thousand-company reference pool, not an exact census or a list of addressable buyers. Assume ten to thirty percent eventually fit the supported variant workflow and can buy, after technical readiness and source coverage are established. That means two hundred to six hundred hypothetical eligible firms. At one paid lab account per eligible firm and $300 monthly, the conditional full-coverage annual subscription pool is $720,000 to $2,160,000; the twenty-percent fit case is $1,440,000. This is a scoped, assumed serviceable segment, not a measured market gap, near-term revenue forecast or obtainable share. The present PTEN-only prototype may qualify far fewer accounts; the fit rate is not measured and commercial readiness is not established. Actual revenue must further multiply by conversion and retention. The first twenty hypothetical lab accounts would represent $72,000 annualized subscriptions, already shown as a commercial scenario; no customers are claimed. Universities and other countries are excluded rather than mixed into this denominator. Global TAM is not established. Pilot evidence and a deduplicated prospect list must replace the assumed fit rate before this is used for fundraising forecasts.''')
for i,(head,num,sub) in enumerate([('REFERENCE POOL','2,000','US biotech firms · rounded'),('ASSUMED FIT','10–30%','200–600 eligible firms'),('PROPOSED PRICE','$3,600','per account per year')]):
 x=.65+i*4.08;txt(s,head,x,2.13,3.8,.3,11,ACC,True);txt(s,num,x,2.8,3.8,.8,39,INK,True);txt(s,sub,x,3.85,3.8,.6,18,MUT)
rect(s,.65,4.85,12.02,1.23,INK,radius=True)
txt(s,'20% fit case: 2,000 × 20% × 1 account × $3,600 = $1.44M/year',.92,5.2,11.5,.7,25,WHITE,True)
txt(s,'Conditional segment value, not forecast revenue. Academic and international expansion excluded.',.65,6.4,12,.42,16,MUT)

#14
s=slide('Charge for a bounded research workspace','Business model hypothesis','All prices, customer counts and revenue values on this slide are proposed assumptions. Arithmetic only: 20 labs × US$300/month × 12 = US$72,000/year; no traction, forecast, profit or valuation claim.',
'''30 seconds. Pricing is a hypothesis to test with the pilot participants. A research individual tier could start around $30 per month and a small-lab tier around $300 per month, with explicit execution budgets and usage limits. These prices are not current commercial offers. The illustrative calculation uses only lab subscriptions: twenty hypothetical labs at three hundred dollars monthly produce seventy-two thousand dollars of annual subscription revenue before costs, churn, discounts or taxes. It is not booked ARR or a financial forecast. The one recorded model-run cost is not enough to infer unit economics; measure support, compute, storage and model usage during pilots.''')
card(s,.65,2.12,5.78,2.25,'Individual researcher','US$30 / month · proposed\nBounded usage; validate with pilots.')
card(s,6.88,2.12,5.79,2.25,'Small laboratory','US$300 / month · proposed\nShared review workflow; bounded usage.')
txt(s,'20 labs × $300 / month × 12',.65,5.0,8.3,.65,32,INK,True)
txt(s,'= $72,000 / year',.65,5.88,8.3,.65,34,COP,True)
pill(s,'ILLUSTRATIVE · NOT TRACTION',9.0,5.55,3.67,INK)
#15
s=slide('Pilot the decision, then earn expansion','Adoption and evaluation','Proposed evaluation plan; pilot counts, durations and success criteria are design hypotheses, not completed customer research.',
'''30 seconds. Propose three to five design-partner labs, starting with agreed public or permissioned nonclinical questions. Freeze the questions and evaluation rubric before the trial. Have a domain expert score applicability, source support and error handling, while a researcher records completion time and time spent checking. Compare against each lab’s current process; do not publish an efficiency gain without the measurements and context. Expand to another gene only after source acquisition, assay interpretation and mapping checks pass. Department adoption should follow evidence of repeat use, exportability and acceptable support costs.''')
for i,(lab,body) in enumerate([('01  Partner','Recruit 3–5 labs.\nAgree on a bounded question.'),('02  Evaluate','Expert review + task success.\nMeasure time and errors.'),('03  Expand','Validate the next gene.\nTest repeat use and lab fit.')]):
 x=.65+i*4.08;card(s,x,2.25,3.83,2.82,lab,body)
txt(s,'What earns the next stage?',.65,5.62,5,.55,27,INK,True)
txt(s,'Traceable results. Useful unknowns.\nResearchers who return.',6.4,5.52,6.27,1.2,24,ACC,True)
#16 close
s=slide('Help shape a research decision worth trusting','The ask','Roadmap and partner request are prospective. Argon status: source/argon-check.json. Gene expansion and independent scientific evaluation remain future milestones.',
'''25 seconds. The next milestones are a verified Argon scientific run, independent researcher evaluation of this bounded workflow, and validated expansion beyond the initial PTEN mission. We are seeking pilot laboratories and scientific advisers who can bring a scoped research question and challenge the evidence, assumptions and next-step usefulness. Longer-term, portable research packages and permissioned collaboration may extend the workflow, but are not required for this initial demonstration. Close on one request: bring one research question and help us test whether the resulting decision record is useful.''')
image(s,'emblem.png',.74,2.05,1.3,1.75)
txt(s,'Bring one research question.',2.65,2.24,9.8,.83,38,INK,True)
txt(s,'Seeking pilot labs and scientific advisers.',2.65,3.32,9.7,.65,26,MUT)
for x,num,label in [(.65,'01','Verify Argon execution'),(4.73,'02','Evaluate with researchers'),(8.81,'03','Validate gene expansion')]:
 line(s,x,5.13,x+3.83,5.13,ACC,2);txt(s,num,x,5.4,3.8,.5,28,COP,True);txt(s,label,x,6.1,3.8,.65,22,INK,True)
# 3D diagrams are editable native objects. Document scientific versus conceptual views.
for idx in [2,5,7,9,13]:
 extra='\n\n3D DIAGRAM: '+('Scientific orthographic projection of actual experimental C-alpha coordinates from source/mapping-audit.json. Rigid view transform only; missing regions are not connected. Side-chain/mutant geometry is not generated.' if idx==5 else 'Conceptual dimensional schematic built from editable native PowerPoint faces, spheres and connectors. Spatial placement is a presentation choice, not measured geometry, runtime motion or a physical simulation.')
 prs.slides[idx].notes_slide.notes_text_frame.text+=extra
 notes[idx]['notes']+=extra
 ledger[idx]['diagram']='source/diagrams_3d.py'+('; source/mapping-audit.json' if idx==5 else '')
(OUT/'source/3d-diagram-manifest.json').write_text(json.dumps({'slides':[3,6,8,10,14],'geometry':'editable native PowerPoint polygons, spheres and connectors','molecule':molecular_diagram,'conceptual_slides':[3,8,10,14]},indent=2))
# Add deliberate slide-level metadata, shape IDs and bounds audit
bounds=[]
for i,s in enumerate(prs.slides,1):
 for sh in s.shapes:
  if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+10 or sh.top+sh.height>prs.slide_height+10:bounds.append(dict(slide=i,name=sh.name))
 if not s.has_notes_slide:raise RuntimeError('Missing notes')
assert not bounds,bounds
prs.save(OUT/'Bramha-Atlas-Pitch.pptx')
(OUT/'speaker-notes.txt').write_text('\n\n'.join(f"SLIDE {n['slide']:02d} — {n['title']}\n{n['notes']}\n\nSources: {n['sources']}" for n in notes))
(OUT/'claim-source-ledger.json').write_text(json.dumps(ledger,indent=2))
(OUT/'source/scene.json').write_text(json.dumps(scene,indent=2))
(OUT/'source/asset-usage.json').write_text(json.dumps(asset_uses,indent=2))
print(f'Built {len(prs.slides)} editable slides with speaker notes; all shape bounds pass.')
