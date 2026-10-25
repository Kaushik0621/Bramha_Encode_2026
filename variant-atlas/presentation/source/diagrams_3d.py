"""Editable dimensional diagrams: oblique 3D solids and an actual-coordinate 3D
protein projection. No generative molecular geometry or runtime events are used.
Every face, connector, sphere and label is a native PowerPoint object.
"""
import math,json
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

BG='F6F3EB';PAPER='FFFDF8';INK='193B3B';ACC='226B60';SEA='A5DFD0';COP='9E5837';DARK='0B1923';WHITE='F2E9D7';MUT='58706F'

def blend(a,b,t):
 return ''.join(f'{round(int(a[i:i+2],16)*(1-t)+int(b[i:i+2],16)*t):02X}' for i in (0,2,4))

def poly(s,pts,color,edge=None):
 # Native editable freeform polygon, in slide coordinates.
 b=s.shapes.build_freeform(pts[0][0]*72,pts[0][1]*72,scale=12700)
 b.add_line_segments([(x*72,y*72) for x,y in pts[1:]],close=True)
 p=b.convert_to_shape();p.fill.solid();p.fill.fore_color.rgb=RGBColor.from_string(color)
 if edge:p.line.color.rgb=RGBColor.from_string(edge);p.line.width=Pt(.5)
 else:p.line.fill.background()
 # Suppress the default Office shape-style shadow; depth comes from geometry.
 p._element.spPr.append(OxmlElement('a:effectLst'))
 return p

def oval(s,x,y,w,h,color,gradient=False):
 p=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(w),Inches(h));p.line.fill.background();p.fill.solid();p.fill.fore_color.rgb=RGBColor.from_string(color)
 p._element.spPr.append(OxmlElement('a:effectLst'))
 if gradient:
  sp=p._element.spPr
  for e in list(sp):
   if e.tag.endswith('solidFill'):sp.remove(e)
  fill=OxmlElement('a:gradFill');stops=OxmlElement('a:gsLst')
  for at,c in [(0,blend(color,'FFFFFF',.65)),(38000,blend(color,'FFFFFF',.18)),(100000,blend(color,DARK,.55))]:
   gs=OxmlElement('a:gs');gs.set('pos',str(at));cc=OxmlElement('a:srgbClr');cc.set('val',c);gs.append(cc);stops.append(gs)
  fill.append(stops);ln=OxmlElement('a:lin');ln.set('ang','5400000');ln.set('scaled','1');fill.append(ln);sp.insert(2,fill)
 return p

def solid(s,x,y,w,h,dx=.38,dy=-.3,color=ACC):
 # Oblique projection of a rectangular prism; all three faces remain editable.
 poly(s,[(x,y),(x+w,y),(x+w+dx,y+dy),(x+dx,y+dy)],blend(color,'FFFFFF',.60))
 poly(s,[(x+w,y),(x+w+dx,y+dy),(x+w+dx,y+h+dy),(x+w,y+h)],blend(color,DARK,.25))
 poly(s,[(x,y),(x+w,y),(x+w,y+h),(x,y+h)],color)

def cylinder(s,x,y,w,h,color=ACC):
 # Native ellipses + faces form a cylinder with clear dimensional lighting.
 e=.28
 poly(s,[(x,y),(x+w,y),(x+w,y+h),(x,y+h)],color)
 oval(s,x,y+h-e/2,w,e,blend(color,DARK,.15))
 oval(s,x,y-e/2,w,e,blend(color,'FFFFFF',.6))

def workflow(s,txt,line,rect,**kw):
 # Floating evidence sheets and raised platforms show the handoff sequence.
 stages=[('Ask','Declare the question'),('Inspect','Check the evidence'),('Investigate','Run scoped work'),('Review','Choose the next step')]
 for i,(label,sub) in enumerate(stages):
  x=.72+i*3.06;y=4.02-i*.16
  oval(s,x-.05,y+.97,2.55,.28,'DDE3D9')
  solid(s,x,y,2.36,.72,.32,-.32,INK if i==3 else ACC)
  # an inset floating sheet on the platform, with a dimensional lip
  solid(s,x+.42,y-.75,1.45,.12,.22,-.18,PAPER)
  if i==0:
   line(s,x+.67,y-.82,x+1.59,y-.82,COP,2)
   line(s,x+.67,y-.61,x+1.35,y-.61,'96ABA4',1.5)
  elif i==1:
   for j in range(3):
    oval(s,x+.64+j*.36,y-.83,.16,.16,COP if j==1 else ACC)
  elif i==2:
   for a,b in [((x+.72,y-.79),(x+1.13,y-.57)),((x+1.13,y-.57),(x+1.6,y-.85))]:line(s,*a,*b,ACC,1.5)
   for xx,yy in [(x+.72,y-.79),(x+1.13,y-.57),(x+1.6,y-.85)]:oval(s,xx-.09,yy-.09,.18,.18,ACC,True)
  else:
   line(s,x+.68,y-.78,x+.94,y-.57,ACC,3);line(s,x+.94,y-.57,x+1.58,y-.98,ACC,3)
  txt(s,f'0{i+1}',x+.18,y+.2,.45,.3,13,SEA,True)
  txt(s,label,x+.73,y+.16,1.55,.4,21,WHITE,True)
  txt(s,sub,x-.02,y+1.38,2.86,.65,18,MUT)
  if i<3:
   line(s,x+2.71,y+.27,x+3.03,y+.11,COP,2.4)
   poly(s,[(x+3.03,y+.11),(x+2.9,y+.08),(x+2.96,y+.24)],COP)
 txt(s,'A checked research brief, with reproducible artifacts and explicit unknowns.',.75,6.32,11.9,.44,22,INK,True)
 txt(s,'3D WORKFLOW SCHEMATIC',.75,2.12,11,.25,10,ACC,True)

def molecule(s,txt,line,rect,source_dir,**kw):
 d=json.loads((source_dir/'mapping-audit.json').read_text())
 mappings=[v for v in d['structure']['mappings'] if v.get('coordinate') and v['status']=='mapped']
 # Rigid rotation and orthographic projection; original coordinates are immutable.
 ca,sa=math.cos(.65),math.sin(.65);cb,sb=math.cos(-.35),math.sin(-.35)
 points=[]
 for v in mappings:
  c=v['coordinate'];x,y,z=c['x'],c['y'],c['z'];xx=ca*x+sa*z;zz=-sa*x+ca*z;yy=cb*y-sb*zz;zz=sb*y+cb*zz
  points.append((v['reference_position'],xx,yy,zz))
 xs=[p[1] for p in points];ys=[p[2] for p in points];zs=[p[3] for p in points]
 scale=min(5.5/(max(xs)-min(xs)),3.63/(max(ys)-min(ys)))
 cx=(max(xs)+min(xs))/2;cy=(max(ys)+min(ys))/2
 projected={p[0]:(4.1+(p[1]-cx)*scale,4.21-(p[2]-cy)*scale,p[3]) for p in points}
 rect(s,.65,1.98,7.25,4.78,DARK,radius=True)
 txt(s,'PTEN · 1D5R · chain A',.91,2.19,5.7,.3,16,SEA,True)
 # Painter's algorithm gives proper visual depth to the reference backbone.
 segments=[]
 for a,b in zip(points,points[1:]):
  pa,pb=projected[a[0]],projected[b[0]]
  # Never join across an unobserved gap in the experimental chain.
  if b[0]-a[0]==1 and math.dist((a[1],a[2],a[3]),(b[1],b[2],b[3]))<5:
   segments.append(((a[3]+b[3])/2,pa,pb))
 for z,a,b in sorted(segments):
  t=(z-min(zs))/(max(zs)-min(zs));color=blend('254250','9ABFB7',t)
  line(s,a[0],a[1],b[0],b[1],color,4.4)
 for pos,x,y,z in sorted(points,key=lambda p:p[3]):
  px,py,_=projected[pos];t=(z-min(zs))/(max(zs)-min(zs));r=.032+.018*t
  oval(s,px-r,py-r,2*r,2*r,blend('2C4B58','B2D5CC',t),True)
 labels={16:('Y16T',1.03,2.85),124:('C124S',5.61,3.1),172:('R172H',5.67,5.75)}
 for pos,(label,lx,ly) in labels.items():
  px,py,_=projected[pos]
  line(s,lx+.4,ly+.18,px,py,COP if pos==124 else SEA,1)
  oval(s,px-.105,py-.105,.21,.21,'DDA476' if pos==124 else SEA,True)
  txt(s,label,lx,ly,1.6,.34,17,'E8C09A' if pos==124 else SEA,True)
 txt(s,'Experimental Cα trace · rigid 3D projection · reference only',.91,6.37,6.7,.25,10,'AEC3C8')
 txt(s,'Identity checked',8.5,2.35,4.1,.4,23,ACC,True)
 txt(s,'Accepted reference-to-chain mapping.',8.5,2.96,4.05,.8,20,MUT)
 txt(s,'Context preserved',8.5,4.04,4.1,.4,23,ACC,True)
 txt(s,'All three positions are mapped.\nThe mutant fold was not computed.',8.5,4.67,4.03,1.3,20,MUT)
 return {'coordinate_source':'source/mapping-audit.json','projection':'rigid rotation + orthographic','atoms':len(points),'gaps':'not connected','highlighted_reference_positions':[16,124,172]}

def agents(s,txt,line,rect,**kw):
 rect(s,.65,2.04,7.55,4.63,DARK,radius=True)
 txt(s,'RECORDED TASK DEPENDENCIES',.91,2.22,6.8,.25,10,SEA,True)
 # Real dependency handoffs, drawn before the nodes; the bus avoids crossing labels.
 for a,b in [((4.0,2.92),(1.62,4.05)),((1.62,4.05),(3.98,4.05)),((1.62,4.05),(1.62,3.55)),((1.62,3.55),(6.36,3.55)),((6.36,3.55),(6.36,4.05)),((3.98,4.05),(2.8,4.05)),((2.8,4.05),(2.8,5.4)),((2.8,5.4),(3.98,5.68)),((6.36,4.05),(3.98,5.68))]:
  line(s,*a,*b,'5F9B97',2)
 def pod(cx,cy,label,sub,color=ACC):
  oval(s,cx-.67,cy+.36,1.6,.22,'07131D')
  solid(s,cx-.6,cy+.13,1.18,.2,.25,-.17,color)
  oval(s,cx-.27,cy-.29,.55,.55,SEA if color==ACC else 'DDA476',True)
  txt(s,label,cx-.9,cy+.61,2.0,.35,18,WHITE,True,align=PP_ALIGN.CENTER)
  txt(s,sub,cx-.9,cy+1.01,2.0,.3,11,'ACC4C7',align=PP_ALIGN.CENTER)
 pod(3.98,2.82,'','',COP)
 txt(s,'Research lead',4.77,2.65,2.8,.35,21,WHITE,True)
 txt(s,'Model orchestrator',4.77,3.04,2.75,.27,12,SEA)
 for cx,lab,sub in [(1.62,'Identity','Deterministic'),(3.98,'Assay','Model + tools'),(6.36,'Structure','Model + tools')]:pod(cx,4.08,lab,sub)
 # Verifier is lower and has labels beside it, avoiding the stage footer.
 solid(s,3.36,5.94,1.18,.18,.25,-.17,ACC);oval(s,3.70,5.52,.55,.55,SEA,True)
 txt(s,'Verifier',4.84,5.68,2.6,.4,20,WHITE,True);txt(s,'Mechanical checks',4.84,6.13,2.6,.3,11,SEA)
 txt(s,'5',8.85,2.32,3.4,.8,46,INK,True);txt(s,'model requests',8.85,3.17,3.7,.4,21,MUT)
 txt(s,'10',8.85,3.94,3.4,.8,46,INK,True);txt(s,'deterministic tool calls',8.85,4.85,3.7,.75,21,MUT)
 txt(s,'Recorded run\nArgon unverified',8.85,5.89,3.7,.75,17,ACC,True)

def architecture(s,txt,line,rect,**kw):
 # Two physically separate stores, a raised computation deck and a human-review
 # layer. The 3D stack encodes responsibilities, not unsupported cloud migration.
 # Riser spines behind the floating planes.
 for x in [2.4,6.7]:line(s,x,5.35,x+1.0,2.7,'B5C9BE',3)
 # Front foundation: separate persistent stores.
 solid(s,.93,5.35,3.35,.63,.45,-.37,INK)
 cylinder(s,1.25,4.68,.92,.62,ACC)
 txt(s,'MongoDB Atlas',2.39,4.68,2.1,.4,18,INK,True)
 txt(s,'Public corpus',1.16,5.54,2.94,.32,18,WHITE,True)
 solid(s,4.79,5.35,3.35,.63,.45,-.37,INK)
 solid(s,5.14,4.62,1.0,.52,.22,-.18,COP)
 txt(s,'Local store',6.37,4.65,2.0,.4,18,INK,True)
 txt(s,'Project state',5.03,5.54,2.9,.32,18,WHITE,True)
 # Scientific computation tier.
 solid(s,2.0,3.86,5.65,.55,.64,-.48,ACC)
 for xx in [2.55,3.12,3.69]:solid(s,xx,3.4,.4,.3,.15,-.12,INK)
 oval(s,6.4,3.18,.56,.56,'DDA476',True)
 txt(s,'Python tools + scoped model reviews',2.25,4.01,5.1,.35,19,WHITE,True)
 # Reviewed output tier. Sheets have thickness and retain editable text.
 solid(s,3.0,2.45,4.67,.48,.54,-.38,PAPER)
 txt(s,'Researcher review',3.26,2.51,4.12,.4,22,INK,True)
 # Clean right-hand explanations remain screen-aligned for readability.
 for yy,head,sub in [(2.14,'REVIEW','Inspect evidence. Decide follow-up.'),(3.77,'COMPUTE','Numerical tools + bounded review.'),(5.39,'PERSIST','Public corpus / local projects:\nseparate stores.')]:
  txt(s,head,9.18,yy,3.45,.25,11,ACC,True);txt(s,sub,9.18,yy+.43,3.5,.92,18,MUT)
 txt(s,'3D architecture schematic · responsibilities and storage boundaries',.93,6.55,11.6,.25,11,MUT)
